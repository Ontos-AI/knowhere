from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://test:test@localhost/test")
os.environ.setdefault("TMP_PATH", "/tmp/knowhere-test")
os.environ.setdefault("S3_BUCKET_NAME", "test-uploads")
os.environ.setdefault("S3_ACCESS_KEY_ID", "test")
os.environ.setdefault("S3_SECRET_ACCESS_KEY", "test")
os.environ.setdefault("S3_TEMP_PATH", "/tmp")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.opc.constants import CONTENT_TYPE as CT  # noqa: E402

from app.services.document_ingestion.page_estimator import PageEstimator  # noqa: E402
from app.services.document_parser.formats.docx.block_stream import (  # noqa: E402
    iter_block_items,
)
from app.services.document_parser.formats.docx.word_package import (  # noqa: E402
    normalize_word_package_bytes,
    open_word_document,
)

_DOCM_MAIN_CONTENT_TYPE = (
    "application/vnd.ms-word.document.macroEnabled.main+xml"
)
_DOTX_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)
_PARAGRAPH_TEXT = "Macro-enabled Word sample"


def test_normalize_word_package_bytes_leaves_standard_docx_unchanged() -> None:
    package_bytes = _build_docx_bytes(_PARAGRAPH_TEXT)

    assert normalize_word_package_bytes(package_bytes) is package_bytes


def test_normalize_word_package_bytes_rewrites_docm_content_type() -> None:
    package_bytes = _with_main_content_type(
        _build_docx_bytes(_PARAGRAPH_TEXT),
        _DOCM_MAIN_CONTENT_TYPE,
    )

    with pytest.raises(ValueError, match="is not a Word file"):
        Document(io.BytesIO(package_bytes))

    normalized = normalize_word_package_bytes(package_bytes)
    document = Document(io.BytesIO(normalized))

    assert _read_main_content_type(normalized) == CT.WML_DOCUMENT_MAIN
    assert document.paragraphs[0].text == _PARAGRAPH_TEXT


def test_iter_block_items_opens_macro_enabled_word_package() -> None:
    package_bytes = _with_main_content_type(
        _build_docx_bytes(_PARAGRAPH_TEXT),
        _DOCM_MAIN_CONTENT_TYPE,
    )

    blocks = list(iter_block_items(package_bytes))
    paragraph_texts = [
        block[1].text if hasattr(block[1], "text") else str(block[1])
        for block in blocks
        if block[2] == "PTXT"
    ]

    assert _PARAGRAPH_TEXT in paragraph_texts


def test_open_word_document_accepts_template_content_type(tmp_path: Path) -> None:
    path = tmp_path / "sample.dotx"
    path.write_bytes(
        _with_main_content_type(
            _build_docx_bytes(_PARAGRAPH_TEXT),
            _DOTX_MAIN_CONTENT_TYPE,
        )
    )

    document = open_word_document(str(path))

    assert document.paragraphs[0].text == _PARAGRAPH_TEXT


def test_page_estimator_counts_macro_enabled_docx(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    path.write_bytes(
        _with_main_content_type(
            _build_docx_bytes(_PARAGRAPH_TEXT),
            _DOCM_MAIN_CONTENT_TYPE,
        )
    )

    estimate = PageEstimator.estimate_workload(str(path))

    assert estimate.used_fallback is False
    assert estimate.method == "docx_words"
    assert estimate.page_count >= 1


def _build_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_main_content_type(package_bytes: bytes, content_type: str) -> bytes:
    source = io.BytesIO(package_bytes)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(output, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    CT.WML_DOCUMENT_MAIN.encode("ascii"),
                    content_type.encode("ascii"),
                )
            dst.writestr(item.filename, data, compress_type=item.compress_type)
    return output.getvalue()


def _read_main_content_type(package_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(package_bytes), "r") as archive:
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
    marker = 'PartName="/word/document.xml" ContentType="'
    start = content_types.index(marker) + len(marker)
    end = content_types.index('"', start)
    return content_types[start:end]
