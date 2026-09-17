from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pytest
from docx import Document

from app.services.document_parser.orchestration.office_container_inspection import (
    CONTENT_TYPES_MEMBER,
    CONTENT_TYPES_NAMESPACE,
)
from shared.core.exceptions.domain_exceptions import ValidationException


def _compat_module():
    from app.services.document_ingestion import office_compat_normalizer

    return office_compat_normalizer

_DOCX_MACRO_CONTENT_TYPE = (
    "application/vnd.ms-word.document.macroEnabled.main+xml"
)
_DOCX_TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)
_XLSX_MACRO_CONTENT_TYPE = (
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
)


def _write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("compat-body")
    document.save(path)


def _write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "compat-body"
    workbook.save(path)


def _rewrite_override_content_type(
    source_path: Path,
    destination_path: Path,
    part_name: str,
    content_type: str,
) -> None:
    override_tag = f"{{{CONTENT_TYPES_NAMESPACE}}}Override"
    with zipfile.ZipFile(source_path, "r") as source_archive:
        with zipfile.ZipFile(destination_path, "w") as destination_archive:
            for info in source_archive.infolist():
                data = source_archive.read(info.filename)
                if info.filename == CONTENT_TYPES_MEMBER:
                    root = ElementTree.fromstring(data)
                    for override in root.findall(override_tag):
                        if override.get("PartName") == part_name:
                            override.set("ContentType", content_type)
                    data = ElementTree.tostring(
                        root,
                        encoding="UTF-8",
                        xml_declaration=True,
                    )
                destination_archive.writestr(info, data)


def test_standard_docx_is_left_unchanged(tmp_path: Path) -> None:
    source_path = tmp_path / "plain.docx"
    _write_docx(source_path)

    result = _compat_module().normalize_office_source(str(source_path))
    assert result.file_path == str(source_path)
    assert result.conversion is None


def test_standard_xlsx_is_left_unchanged(tmp_path: Path) -> None:
    source_path = tmp_path / "plain.xlsx"
    _write_xlsx(source_path)

    result = _compat_module().normalize_office_source(str(source_path))
    assert result.file_path == str(source_path)
    assert result.conversion is None


def test_non_office_extension_is_left_unchanged(tmp_path: Path) -> None:
    source_path = tmp_path / "notes.pdf"
    source_path.write_bytes(b"%PDF")

    result = _compat_module().normalize_office_source(str(source_path))
    assert result.file_path == str(source_path)
    assert result.conversion is None


def test_unreadable_zip_docx_is_left_unchanged(tmp_path: Path) -> None:
    source_path = tmp_path / "broken.docx"
    source_path.write_bytes(b"this is not a docx package")

    result = _compat_module().normalize_office_source(str(source_path))
    assert result.file_path == str(source_path)
    assert result.conversion is None


def test_unknown_content_type_is_left_unchanged(tmp_path: Path) -> None:
    source_path = tmp_path / "plain.docx"
    _write_docx(source_path)
    variant_path = tmp_path / "unknown.docx"
    _rewrite_override_content_type(
        source_path,
        variant_path,
        "/word/document.xml",
        "application/octet-stream",
    )

    result = _compat_module().normalize_office_source(str(variant_path))
    assert result.file_path == str(variant_path)
    assert result.conversion is None


def test_macro_docx_is_converted(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_path = tmp_path / "plain.docx"
    _write_docx(source_path)
    variant_path = tmp_path / "macro.docx"
    _rewrite_override_content_type(
        source_path,
        variant_path,
        "/word/document.xml",
        _DOCX_MACRO_CONTENT_TYPE,
    )
    converted_path = tmp_path / "converted.docx"
    _write_docx(converted_path)

    def fake_normalize(path: str, outdir: str = ".") -> tuple[str, str]:
        assert path == str(variant_path)
        return str(converted_path), converted_path.name

    monkeypatch.setattr(_compat_module(), "normalize_docx_variant", fake_normalize)

    result = _compat_module().normalize_office_source(str(variant_path))
    assert result.file_path == str(converted_path)
    assert result.conversion == {"content_type": _DOCX_MACRO_CONTENT_TYPE}


def test_template_docx_is_converted(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_path = tmp_path / "plain.docx"
    _write_docx(source_path)
    variant_path = tmp_path / "template.docx"
    _rewrite_override_content_type(
        source_path,
        variant_path,
        "/word/document.xml",
        _DOCX_TEMPLATE_CONTENT_TYPE,
    )
    converted_path = tmp_path / "converted.docx"
    _write_docx(converted_path)

    monkeypatch.setattr(
        _compat_module(),
        "normalize_docx_variant",
        lambda path, outdir=".": (str(converted_path), converted_path.name),
    )

    result = _compat_module().normalize_office_source(str(variant_path))
    assert result.file_path == str(converted_path)
    assert result.conversion == {"content_type": _DOCX_TEMPLATE_CONTENT_TYPE}


def test_macro_xlsx_is_converted(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_path = tmp_path / "plain.xlsx"
    _write_xlsx(source_path)
    variant_path = tmp_path / "macro.xlsx"
    _rewrite_override_content_type(
        source_path,
        variant_path,
        "/xl/workbook.xml",
        _XLSX_MACRO_CONTENT_TYPE,
    )
    converted_path = tmp_path / "converted.xlsx"
    _write_xlsx(converted_path)

    monkeypatch.setattr(
        _compat_module(),
        "normalize_xlsx_variant",
        lambda path, outdir=".": (str(converted_path), converted_path.name),
    )

    result = _compat_module().normalize_office_source(str(variant_path))
    assert result.file_path == str(converted_path)
    assert result.conversion == {"content_type": _XLSX_MACRO_CONTENT_TYPE}


def test_ole_container_is_rejected_as_encrypted(tmp_path: Path) -> None:
    source_path = tmp_path / "locked.docx"
    source_path.write_bytes(_compat_module().OLE_CFBF_SIGNATURE + b"\x00")

    with pytest.raises(ValidationException) as raised:
        _compat_module().normalize_office_source(str(source_path))

    assert raised.value.user_message == (
        "Invalid file: the uploaded .docx file is encrypted or "
        "password-protected. Please unlock the file and upload again."
    )
    assert raised.value.violations == [
        {
            "field": "file",
            "description": "Encrypted or password-protected Office file",
        }
    ]


def test_ole_xlsx_container_is_rejected_as_encrypted(tmp_path: Path) -> None:
    source_path = tmp_path / "locked.xlsx"
    source_path.write_bytes(_compat_module().OLE_CFBF_SIGNATURE + b"\x00")

    with pytest.raises(ValidationException) as raised:
        _compat_module().normalize_office_source(str(source_path))

    assert ".xlsx" in raised.value.user_message
