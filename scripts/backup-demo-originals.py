"""Back up authoritative S3 originals and fail closed on missing/invalid files."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import boto3
import fitz
from docx import Document


def validate_original(path: Path) -> dict[str, object]:
    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as document:
            if document.is_encrypted or document.page_count < 1:
                raise ValueError(
                    "PDF cannot be opened without a password or has no pages"
                )
            for page in document:
                page.get_text()
            return {"format": "pdf", "pages": document.page_count, "is_valid": True}
    if path.suffix.lower() == ".docx":
        with ZipFile(path) as archive:
            if (
                archive.testzip() is not None
                or "word/document.xml" not in archive.namelist()
            ):
                raise ValueError("Invalid DOCX archive")
        document = Document(str(path))
        return {
            "format": "docx",
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "is_valid": True,
        }
    raise ValueError("Original format is not supported by the backup validator")


def hash_original(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as original:
        for block in iter(lambda: original.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def back_up_originals(
    *, directory: Path, sources: list[dict[str, Any]], profile: str, bucket: str
) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    client = boto3.Session(profile_name=profile).client("s3", region_name="us-west-1")
    inventoryPath: Path = directory / "inventory.json"
    priorInventory: list[dict[str, Any]] = (
        json.loads(inventoryPath.read_text()) if inventoryPath.exists() else []
    )
    inventory: list[dict[str, Any]] = []
    for source in sources:
        sourceId: str = str(source["source_id"])
        fileName: str = str(source["file_name"])
        if (
            Path(fileName).name != fileName
            or "/" in sourceId
            or "\\" in sourceId
            or sourceId in (".", "..")
        ):
            raise ValueError("Unsafe backup source path")
        path: Path = directory / sourceId / fileName
        path.parent.mkdir(parents=True, exist_ok=True)
        entry: dict[str, Any] = {
            **source,
            "relative_path": path.relative_to(directory).as_posix(),
        }
        try:
            head = client.head_object(Bucket=bucket, Key=source["s3_key"])
            size: int = int(head["ContentLength"])
            etag: str = str(head["ETag"])
            prior = next(
                (item for item in priorInventory if item["source_id"] == sourceId), {}
            )
            if (
                not path.exists()
                or path.stat().st_size != size
                or prior.get("etag") != etag
            ):
                temporary: Path = path.with_name(path.name + ".partial")
                client.download_file(bucket, source["s3_key"], str(temporary))
                if temporary.stat().st_size != size:
                    raise ValueError("Downloaded original size does not match S3 HEAD")
                temporary.replace(path)
            entry.update(
                {
                    "size_bytes": size,
                    "sha256": hash_original(path),
                    "etag": etag,
                    "validation": validate_original(path),
                    "status": "verified",
                }
            )
            print(f"{sourceId}: verified {size} bytes", flush=True)
        except Exception as error:
            entry.update({"status": "failed", "error_type": type(error).__name__})
            print(f"{sourceId}: failed ({type(error).__name__})", flush=True)
        inventory.append(entry)
        temporaryInventory = inventoryPath.with_suffix(".partial")
        temporaryInventory.write_text(json.dumps(inventory, indent=2) + "\n")
        temporaryInventory.replace(inventoryPath)
    if len(inventory) != len(sources) or any(
        item["status"] != "verified" for item in inventory
    ):
        raise RuntimeError("Backup is incomplete; do not switch demo publication")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--directory",
        type=Path,
        default=Path.home()
        / "knowhere-demo-originals"
        / datetime.now(timezone.utc).strftime("%Y%m%d"),
    )
    parser.add_argument(
        "--sources",
        type=Path,
        default=Path(__file__).with_name("demo-original-sources.json"),
    )
    parser.add_argument("--profile", default="knowhere")
    parser.add_argument("--bucket", default="knowhere-storage-staging")
    arguments = parser.parse_args()
    back_up_originals(
        directory=arguments.directory,
        sources=json.loads(arguments.sources.read_text()),
        profile=arguments.profile,
        bucket=arguments.bucket,
    )


if __name__ == "__main__":
    main()
